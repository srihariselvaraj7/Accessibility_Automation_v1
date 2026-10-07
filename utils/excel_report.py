import os

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


def generate_excel_report(
    page_url,
    results,
    heading_result,
    image_result,
    button_result,
    link_result,
    contrast_result
):

    os.makedirs("reports", exist_ok=True)

    report_path = os.path.join(
        "reports",
        "accessibility_report.xlsx"
    )

    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "Accessibility Report"

    # ==================================================
    # TITLE
    # ==================================================

    sheet["A1"] = "Accessibility Test Report"

    sheet["A1"].font = Font(
        bold=True,
        size=16
    )

    sheet.merge_cells("A1:I1")

    # ==================================================
    # PAGE INFORMATION
    # ==================================================

    sheet["A3"] = "Page URL"
    sheet["B3"] = page_url

    # ==================================================
    # OVERALL SCAN SUMMARY
    # ==================================================

    sheet["A5"] = "Overall Scan Summary"

    sheet["A5"].font = Font(
        bold=True,
        size=12
    )

    summary_headers = [
        "Scanner",
        "Total",
        "Passed",
        "Failed"
    ]

    summary_header_row = 6

    for column, header in enumerate(
        summary_headers,
        start=1
    ):

        cell = sheet.cell(
            row=summary_header_row,
            column=column
        )

        cell.value = header

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # ==================================================
    # SCANNER RESULTS
    # ==================================================

    scanner_results = [
        ("Headings", heading_result),
        ("Images", image_result),
        ("Buttons", button_result),
        ("Links", link_result),
        ("Color Contrast", contrast_result)
    ]

    current_summary_row = 7

    total_checked = 0
    total_passed = 0
    total_failed = 0

    for scanner_name, result in scanner_results:

        checked = result.checked
        passed = result.passed
        failed = checked - passed

        sheet.cell(
            row=current_summary_row,
            column=1
        ).value = scanner_name

        sheet.cell(
            row=current_summary_row,
            column=2
        ).value = checked

        sheet.cell(
            row=current_summary_row,
            column=3
        ).value = passed

        sheet.cell(
            row=current_summary_row,
            column=4
        ).value = failed

        total_checked += checked
        total_passed += passed
        total_failed += failed

        current_summary_row += 1

    # ==================================================
    # TOTAL
    # ==================================================

    sheet.cell(
        row=current_summary_row,
        column=1
    ).value = "TOTAL"

    sheet.cell(
        row=current_summary_row,
        column=2
    ).value = total_checked

    sheet.cell(
        row=current_summary_row,
        column=3
    ).value = total_passed

    sheet.cell(
        row=current_summary_row,
        column=4
    ).value = total_failed

    for column in range(1, 5):

        sheet.cell(
            row=current_summary_row,
            column=column
        ).font = Font(
            bold=True
        )

    # ==================================================
    # ACCESSIBILITY ISSUES
    # ==================================================

    issue_summary_row = current_summary_row + 3

    sheet.cell(
        row=issue_summary_row,
        column=1
    ).value = "Accessibility Issues Found"

    sheet.cell(
        row=issue_summary_row,
        column=1
    ).font = Font(
        bold=True,
        size=12
    )

    sheet.cell(
        row=issue_summary_row,
        column=2
    ).value = len(results)

    # ==================================================
    # ISSUE DETAILS
    # ==================================================

    header_row = issue_summary_row + 2

    sheet.cell(
        row=header_row,
        column=1
    ).value = "Issue Details"

    sheet.cell(
        row=header_row,
        column=1
    ).font = Font(
        bold=True,
        size=12
    )

    table_header_row = header_row + 1

    headers = [
        "No.",
        "Page URL",
        "Issue Type",
        "XPath",
        "Description",
        "Severity",
        "Priority",
        "Test Case",
        "Screenshot"
    ]

    for column, header in enumerate(
        headers,
        start=1
    ):

        cell = sheet.cell(
            row=table_header_row,
            column=column
        )

        cell.value = header

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # ==================================================
    # ISSUE ROWS
    # ==================================================

    current_row = table_header_row + 1

    for index, result in enumerate(
        results,
        start=1
    ):

        # ------------------------------------------------
        # No.
        # ------------------------------------------------

        sheet.cell(
            row=current_row,
            column=1
        ).value = index

        # ------------------------------------------------
        # Page URL
        # ------------------------------------------------

        sheet.cell(
            row=current_row,
            column=2
        ).value = result.page_url

        # ------------------------------------------------
        # Issue Type
        # ------------------------------------------------

        sheet.cell(
            row=current_row,
            column=3
        ).value = getattr(
            result,
            "issue_type",
            "-"
        )

        # ------------------------------------------------
        # XPath
        # ------------------------------------------------

        sheet.cell(
            row=current_row,
            column=4
        ).value = result.xpath

        # ------------------------------------------------
        # Description
        # ------------------------------------------------

        sheet.cell(
            row=current_row,
            column=5
        ).value = result.description

        # ------------------------------------------------
        # Severity
        # ------------------------------------------------

        sheet.cell(
            row=current_row,
            column=6
        ).value = getattr(
            result,
            "severity",
            "-"
        )

        # ------------------------------------------------
        # Priority
        # ------------------------------------------------

        sheet.cell(
            row=current_row,
            column=7
        ).value = getattr(
            result,
            "priority",
            "-"
        )

        # ------------------------------------------------
        # Test Case
        # ------------------------------------------------

        sheet.cell(
            row=current_row,
            column=8
        ).value = result.test_case

        # ------------------------------------------------
        # Screenshot
        # ------------------------------------------------

        screenshot = getattr(
            result,
            "screenshot",
            None
        )

        screenshot_status = getattr(
            result,
            "screenshot_status",
            "Not captured"
        )

        if (
            screenshot
            and os.path.exists(screenshot)
        ):

            try:

                image = ExcelImage(
                    screenshot
                )

                image.width = 200
                image.height = 120

                sheet.add_image(
                    image,
                    f"I{current_row}"
                )

                sheet.row_dimensions[
                    current_row
                ].height = 100

            except Exception:

                sheet.cell(
                    row=current_row,
                    column=9
                ).value = (
                    "Screenshot unavailable"
                )

        else:

            sheet.cell(
                row=current_row,
                column=9
            ).value = screenshot_status

        current_row += 1

    # ==================================================
    # COLUMN WIDTHS
    # ==================================================

    widths = {
        1: 12,
        2: 45,
        3: 30,
        4: 60,
        5: 70,
        6: 12,
        7: 12,
        8: 15,
        9: 25
    }

    for column, width in widths.items():

        sheet.column_dimensions[
            get_column_letter(column)
        ].width = width

    # ==================================================
    # TEXT ALIGNMENT
    # ==================================================

    for row in sheet.iter_rows(
        min_row=1,
        max_row=current_row - 1
    ):

        for cell in row:

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

    # ==================================================
    # IMPORTANT:
    # NO FREEZE PANES
    # ==================================================

    # Do NOT use:
    #
    # sheet.freeze_panes = "A11"
    #
    # This allows the summary to scroll away
    # when the user scrolls down to issue details.

    # ==================================================
    # SAVE
    # ==================================================

    workbook.save(
        report_path
    )

    return report_path